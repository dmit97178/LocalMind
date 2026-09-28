"""Bounded extraction of local workspace documents."""

from __future__ import annotations

import importlib.util
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from src.core.schemas import PreflightReport, ToolContext, ToolDef, ToolError
from src.tools.file_tools import (
    _file_error,
    _preflight_target,
    _require_readable_file,
    resolve_workspace_target,
)

MAX_SEGMENT_LINES = 200
_DEFAULTS = {
    "max_file_bytes": 32 * 1024 * 1024,
    "max_extracted_chars": 500_000,
    "max_segments_total": 20_000,
    "max_returned_segments": 200,
    "max_returned_chars": 40_000,
    "max_zip_entries": 2_000,
    "max_uncompressed_bytes": 256 * 1024 * 1024,
}
_MODULES = {"pdf": "pypdf", "docx": "docx", "xlsx": "openpyxl", "pptx": "pptx"}
_ZIP_MEMBER = {
    "docx": "word/document.xml",
    "xlsx": "xl/workbook.xml",
    "pptx": "ppt/presentation.xml",
}

DOCUMENT_TOOL_DEF = ToolDef(
    name="document.read",
    description=(
        "读取工作区内一份文档的文本内容并返回带来源定位（页码/段落/表格行/工作表行/幻灯片）的结构化结果。"
        "支持 pdf、docx、xlsx、pptx、txt、md。"
        "范围：DOCX 正文段落与表格（不含页眉页脚）；XLSX 各工作表非空行（公式单元格返回公式文本）；"
        "PPTX 各幻灯片的文本与表格（不含备注）。"
        "不支持 OCR 与扫描型 PDF，不支持 .doc/.xls/.ppt。"
        "结果有长度上限；超过时会在结果中明确说明未返回的部分。"
    ),
    parameters={
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
    },
    risk_level="sensitive",
    result_delivery="model_summary",
)


def make_document_tool(settings: dict):
    from src.tools.registry import RegisteredTool

    limits = {**_DEFAULTS, **settings}

    def handler(params: dict, context: ToolContext) -> dict:
        return read_document(params, context, limits)

    def preflight(params: dict, context: ToolContext) -> PreflightReport:
        return preflight_document(params, context, limits)

    return RegisteredTool(DOCUMENT_TOOL_DEF, handler, preflight)


def preflight_document(
    params: dict, context: ToolContext, limits: dict[str, int]
) -> PreflightReport:
    target = _preflight_target(params, context)
    if isinstance(target, PreflightReport):
        return target
    _require_readable_file(target)
    try:
        size = target.stat().st_size
    except OSError as exc:
        raise _file_error("stat_failed", "文件属性读取失败。") from exc
    if size > limits["max_file_bytes"]:
        raise _file_error("target_too_large", "目标文件超过文档读取大小限制。")
    return PreflightReport(
        effective_risk=context.tool_def.risk_level,
        requires_confirmation=True,
        target_exists=True,
        target_key=str(target),
        approval_profile=context.tool_def.approval_profile,
    )


def read_document(params: dict, context: ToolContext, limits: dict[str, int]) -> dict:
    path = resolve_workspace_target(params.get("path"), context.workspace_root)
    _require_readable_file(path)
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise _file_error("stat_failed", "文件属性读取失败。") from exc
    if size > limits["max_file_bytes"]:
        raise _file_error("target_too_large", "目标文件超过文档读取大小限制。")
    declared = path.suffix.lower().lstrip(".")
    if declared in {"doc", "xls", "ppt"}:
        raise _document_error(
            "document_unsupported_format",
            "该文档格式暂不支持；旧版 Office 文件请先转换为 OOXML 格式。",
            False,
        )
    if declared not in {"pdf", "docx", "xlsx", "pptx", "txt", "md"}:
        raise _document_error(
            "document_unsupported_format",
            "该文件扩展名不在支持的文档格式范围内。",
            False,
        )
    try:
        _validate_container(path, declared, limits)
        if declared in _MODULES and importlib.util.find_spec(_MODULES[declared]) is None:
            raise _document_error(
                "parser_unavailable",
                f"读取 .{declared} 需要的解析组件不可用，请重新安装依赖。",
                True,
                format=declared,
            )
        extraction_result, extraction = _extract(path, declared, limits)
    except ToolError:
        raise
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise _document_error(
            "document_corrupt", "文件已损坏或格式不完整，无法读取。", False,
            exception_type=type(exc).__name__,
        ) from exc
    except Exception as exc:
        raise _document_error(
            "document_extract_failed", "文档文本提取失败，请重试。", True,
            parser=_parser_name(declared), exception_type=type(exc).__name__,
        ) from exc

    return _bounded_result(
        path,
        context.workspace_root,
        declared,
        size,
        extraction_result.segments,
        extraction,
        limits,
        chars_extracted=extraction_result.chars,
        extraction_reason=extraction_result.reason,
    )


def _validate_container(path: Path, declared: str, limits: dict[str, int]) -> None:
    with path.open("rb") as stream:
        head = stream.read(8192)
    if declared == "pdf":
        if not head.startswith(b"%PDF-"):
            raise _document_error(
                "document_format_mismatch", "文件内容与扩展名不一致。", False
            )
        return
    if declared in {"txt", "md"}:
        if b"\0" in head:
            raise _document_error(
                "document_format_mismatch", "文件内容与扩展名不一致。", False
            )
        return
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        raise _document_error(
            "document_encrypted",
            "该文件是加密的 Office 文档或旧版二进制格式容器，首版不支持。",
            False,
        )
    if not head.startswith(b"PK\x03\x04"):
        raise _document_error(
            "document_format_mismatch", "文件内容与扩展名不一致。", False
        )
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if _ZIP_MEMBER[declared] not in {item.filename for item in entries}:
                raise _document_error(
                    "document_format_mismatch", "文件内容与扩展名不一致。", False
                )
            if len(entries) > limits["max_zip_entries"] or sum(
                item.file_size for item in entries
            ) > limits["max_uncompressed_bytes"]:
                raise _document_error(
                    "document_expansion_too_large", "文档展开后超过安全限制。", False
                )
    except zipfile.BadZipFile as exc:
        raise _document_error(
            "document_corrupt", "文件已损坏或格式不完整，无法读取。", False
        ) from exc


@dataclass
class _Extraction:
    limits: dict[str, int]
    segments: list[dict[str, object]] = field(default_factory=list)
    chars: int = 0
    reason: str | None = None

    def add(self, segment: dict[str, object]) -> bool:
        self.segments.append(segment)
        self.chars += len(str(segment["text"]))
        if self.chars > self.limits["max_extracted_chars"]:
            self.reason = "max_extracted_chars"
        elif len(self.segments) > self.limits["max_segments_total"]:
            self.reason = "max_segments_total"
        return self.reason is None


def _extract(
    path: Path, declared: str, limits: dict[str, int]
) -> tuple[_Extraction, str]:
    result = _Extraction(limits)
    if declared in {"txt", "md"}:
        text = path.read_text(encoding="utf-8")
        lines = text.splitlines()
        for index in range(0, len(lines), MAX_SEGMENT_LINES):
            if not result.add({
                "locator": {"kind": "line_range", "start": index + 1,
                            "end": min(index + MAX_SEGMENT_LINES, len(lines))},
                "text": "\n".join(lines[index:index + MAX_SEGMENT_LINES]),
            }):
                break
        return result, "ok" if result.segments else "empty"
    if declared == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise _document_error("document_encrypted", "该 PDF 已加密，无法读取。", False)
        for index, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text and not result.add(
                {"locator": {"kind": "page", "index": index}, "text": text}
            ):
                break
        return result, "ok" if result.segments else (
            "no_text_layer" if len(reader.pages) else "empty"
        )
    if declared == "docx":
        from docx import Document

        document = Document(str(path))
        stop = False
        for index, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip() and not result.add(
                {"locator": {"kind": "paragraph", "index": index},
                 "text": paragraph.text}
            ):
                stop = True
                break
        for table_index, table in enumerate(document.tables if not stop else (), 1):
            for row_index, row in enumerate(table.rows, 1):
                text = "\t".join(cell.text for cell in row.cells)
                if text.strip() and not result.add({
                    "locator": {"kind": "table_row", "index": table_index,
                                "row": row_index}, "text": text
                }):
                    stop = True
                    break
            if stop:
                break
        return result, "ok" if result.segments else "empty"
    if declared == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=False, keep_links=False)
        try:
            stop = False
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if not any(value is not None and str(value) != "" for value in row):
                        continue
                    if not result.add({
                        "locator": {"kind": "sheet", "name": sheet.title,
                                    "row": row_number},
                        "text": "\t".join("" if value is None else str(value)
                                          for value in row),
                    }):
                        stop = True
                        break
                if stop:
                    break
            return result, "ok" if result.segments else "empty"
        finally:
            workbook.close()
    from pptx import Presentation

    presentation = Presentation(str(path))
    for index, slide in enumerate(presentation.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                parts.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
        if parts:
            if not result.add({"locator": {"kind": "slide", "index": index},
                               "text": "\n".join(parts)}):
                break
    return result, "ok" if result.segments else "empty"


def _bounded_result(
    path: Path,
    root: Path,
    declared: str,
    size: int,
    raw_segments: Iterable[dict[str, object]],
    extraction: str,
    limits: dict[str, int],
    *,
    chars_extracted: int,
    extraction_reason: str | None,
) -> dict:
    extracted = list(raw_segments)

    returned: list[dict[str, object]] = []
    chars_returned = 0
    return_reason: str | None = None
    segment_truncated = False
    if extraction == "ok":
        for segment in extracted:
            text = str(segment["text"])
            if len(returned) >= limits["max_returned_segments"]:
                return_reason = "max_returned_segments"
                break
            if chars_returned + len(text) > limits["max_returned_chars"]:
                return_reason = "max_returned_chars"
                if not returned:
                    segment = {**segment, "text": text[:limits["max_returned_chars"]]}
                    returned.append(segment)
                    chars_returned = len(str(segment["text"]))
                    segment_truncated = True
                break
            returned.append(segment)
            chars_returned += len(text)

    reason = extraction_reason or return_reason
    total = None if extraction_reason else len(extracted)
    if extraction == "empty":
        notice = "该文档不含任何文本内容。"
    elif extraction == "no_text_layer":
        notice = "该文档未包含可提取的文本层（可能是扫描件或纯图片文档）；本版本不支持 OCR。"
    elif extraction_reason:
        notice = (
            f"已提取 {len(extracted)} 个文本段、{chars_extracted} 个字符后达到限制；"
            "文档总数未知。"
        )
    elif return_reason:
        remaining_segments = total - len(returned) if total is not None else 0
        notice = (
            f"文档共 {total} 个文本段，本次返回 {len(returned)} 段；"
            f"剩余 {remaining_segments} 段未返回。"
        )
    else:
        notice = None
    return {
        "path": path.relative_to(root).as_posix(), "declared_format": declared,
        "format": declared, "size_bytes": size, "extraction": extraction,
        "segments": returned, "segments_total": total,
        "segments_returned": len(returned), "chars_extracted": chars_extracted,
        "chars_returned": chars_returned, "segment_truncated": segment_truncated,
        "truncated": reason is not None, "truncation_reason": reason, "notice": notice,
    }


def _document_error(
    reason: str, user_message: str, recoverable: bool, **details: object
) -> ToolError:
    return ToolError(
        code="TOOL_EXEC_ERROR", message=f"Document read failed: {reason}",
        user_message=user_message, recoverable=recoverable,
        details={"reason": reason, **details},
    )


def _parser_name(declared: str) -> str:
    return _MODULES.get(declared, "text")
